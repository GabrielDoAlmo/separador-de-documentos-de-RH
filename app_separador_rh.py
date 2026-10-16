import os
import re
import sys
import time
import unicodedata
import threading
import tempfile
import subprocess
import customtkinter as ctk
from tkinter import filedialog, messagebox
from pypdf import PdfReader, PdfWriter

try:
    import fitz  # PyMuPDF
    from PIL import Image
    FOTO_DISPONIVEL = True
except ImportError:
    FOTO_DISPONIVEL = False

# --- CONFIGURAÇÃO PORTÁTIL DO TESSERACT-OCR ---
# Descobre se está rodando como script normal ou como executável (.exe) do PyInstaller
if getattr(sys, 'frozen', False):
    caminho_base = sys._MEIPASS
else:
    caminho_base = os.path.dirname(os.path.abspath(__file__))

# Aponta automaticamente para a pasta 'Tesseract-OCR' localizada na mesma pasta deste app
CAMINHO_TESSERACT = os.path.join(caminho_base, "Tesseract-OCR", "tesseract.exe")

# Variáveis Globais para o sistema de pausa e escolha de adiantamentos
escolha_global = -1
evento_espera = threading.Event()


def normalizar(texto):
    if not texto:
        return ""
    nfkd = unicodedata.normalize("NFKD", texto)
    return "".join([c for c in nfkd if not unicodedata.combining(c)]).upper().strip()


def extrair_texto_ou_ocr(pagina, caminho_pdf, numero_pagina=0):
    """
    MODO DE EMERGÊNCIA: Tenta extrair o texto digital normalmente.
    Se vier vazio (PDF escaneado/imagem pura), ativa o OCR do Tesseract portátil automaticamente.
    """
    texto = pagina.extract_text()
    if texto and len(texto.strip()) > 15:
        return texto.upper()
    
    try:
        doc = fitz.open(caminho_pdf)
        pagina_fitz = doc.load_page(numero_pagina)
        pix = pagina_fitz.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        
        pasta_temp = tempfile.gettempdir()
        caminho_img_temp = os.path.join(pasta_temp, "temp_ocr_emergencia.png")
        img.save(caminho_img_temp)
        doc.close()
        
        comando = [CAMINHO_TESSERACT, caminho_img_temp, "stdout", "-l", "por"]
        resultado = subprocess.run(comando, capture_output=True, text=True, encoding='utf-8')
        texto_ocr = resultado.stdout.upper()
        
        if os.path.exists(caminho_img_temp):
            os.remove(caminho_img_temp)
            
        return texto_ocr
    except Exception as e:
        print(f"[ERRO OCR Emergência]: {e}")
        return ""


def carregar_nomes_de_planilha(caminho_arquivo):
    """Lê uma tabela Excel (.xlsx) ou CSV e extrai a lista de nomes dos funcionários."""
    leitor = []
    try:
        if caminho_arquivo.endswith('.xlsx'):
            from openpyxl import load_workbook
            ws = load_workbook(caminho_arquivo, data_only=True).active
            for row in ws.iter_rows(values_only=True):
                leitor.append([str(c).strip() if c is not None else "" for c in row])
        elif caminho_arquivo.endswith('.csv'):
            import csv, io
            try:
                with open(caminho_arquivo, mode='r', encoding='utf-8', errors='ignore') as f:
                    cont = f.read()
            except Exception:
                with open(caminho_arquivo, mode='r', encoding='latin-1', errors='ignore') as f:
                    cont = f.read()
            if cont:
                delim = ';' if cont.count(';') > cont.count(',') else ','
                leitor = list(csv.reader(io.StringIO(cont), delimiter=delim))
        
        if not leitor:
            return []
        
        col_nome = 0
        header_idx = -1
        for i, row in enumerate(leitor[:10]):
            for c_idx, val in enumerate([str(c).upper().strip() for c in row]):
                if "NOME" in val:
                    col_nome = c_idx
                    header_idx = i
                    break
            if header_idx != -1:
                break
                
        start_row = header_idx + 1 if header_idx != -1 else 0
        nomes = []
        for linha in leitor[start_row:]:
            if len(linha) > col_nome:
                nome = str(linha[col_nome]).strip()
                if nome and "TOTAL" not in nome.upper() and len(nome) > 2:
                    if nome.upper() not in ["NOME", "NOME COMPLETO", "FUNCIONÁRIO", "FUNCIONARIO"]:
                        nome_limpo = " ".join(nome.upper().split())
                        if nome_limpo not in nomes:
                            nomes.append(nome_limpo)
        return nomes
    except Exception as e:
        messagebox.showerror("Erro ao Ler Planilha", f"Não foi possível processar a planilha:\n{e}")
        return []


class RedirecionadorConsole:
    def __init__(self, textbox):
        self.textbox = textbox

    def write(self, texto):
        self.textbox.insert(ctk.END, texto)
        self.textbox.see(ctk.END)

    def flush(self):
        pass


def extrair_todos_valores(caminho_pdf):
    try:
        reader = PdfReader(caminho_pdf)
        texto = normalizar(extrair_texto_ou_ocr(reader.pages[0], caminho_pdf, 0))
        numeros = re.findall(r"\d{1,3}(?:\.\d{3})*,\d{2}", texto)
        return [n for n in numeros if n not in ("0,00", "00,00")]
    except Exception:
        return []


def extrair_valor_linha_adiantamento(caminho_recibo):
    try:
        reader = PdfReader(caminho_recibo)
        texto = normalizar(extrair_texto_ou_ocr(reader.pages[0], caminho_recibo, 0))
        for linha in texto.replace("\r", "\n").split("\n"):
            if "ADIANTAMENTO" in linha:
                numeros = re.findall(r"\d{1,3}(?:\.\d{3})*,\d{2}", linha)
                if numeros:
                    return numeros[-1]
        return None
    except Exception:
        return None


def abrir_janela_escolha(nome, opcoes, caminho_recibo, app_principal):
    janela = ctk.CTkToplevel(app_principal)
    janela.title("Ação Necessária: Decisão do Gestor")
    janela.state("zoomed")
    janela.attributes("-topmost", True)
    janela.grab_set()

    janela.grid_columnconfigure(0, weight=1)
    janela.grid_columnconfigure(1, weight=1)
    janela.grid_rowconfigure(0, weight=1)

    frame_esq = ctk.CTkFrame(janela)
    frame_esq.grid(row=0, column=0, padx=15, pady=15, sticky="nsew")
    
    ctk.CTkLabel(frame_esq, text="Pré-visualização Ampliada do Contracheque:", font=("Arial", 14, "bold")).pack(pady=10)

    if FOTO_DISPONIVEL:
        try:
            with fitz.open(caminho_recibo) as doc:
                page = doc.load_page(0)
                pix = page.get_pixmap(matrix=fitz.Matrix(2.2, 2.2))
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            img.thumbnail((700, 750))
            
            ctk_img = ctk.CTkImage(light_image=img, dark_image=img, size=(img.width, img.height))
            ctk.CTkLabel(frame_esq, image=ctk_img, text="").pack(pady=5, expand=True)
        except Exception:
            ctk.CTkLabel(frame_esq, text="[Erro ao carregar a foto ampliada]").pack(pady=50)
    else:
        ctk.CTkLabel(frame_esq, text="Módulo de imagem indisponível.").pack(pady=50)

    ctk.CTkButton(
        frame_esq, text="🔍 Abrir PDF Completo Original", fg_color="#4B0082", hover_color="#290047",
        height=40, font=("Arial", 13, "bold"), command=lambda: os.startfile(caminho_recibo)
    ).pack(pady=15)

    frame_dir = ctk.CTkFrame(janela, fg_color="transparent")
    frame_dir.grid(row=0, column=1, padx=15, pady=15, sticky="nsew")
    
    ctk.CTkLabel(
        frame_dir, text=f"O sistema esbarrou no ADIANTAMENTO de:\n{nome}\n\nEscolha o SALÁRIO correto abaixo:",
        font=("Arial", 18, "bold"), text_color="#FFA500"
    ).pack(pady=20)
    
    var_opcao = ctk.IntVar(value=0)
    frame_opcoes = ctk.CTkScrollableFrame(frame_dir, width=500, height=350)
    frame_opcoes.pack(pady=10, expand=True)
    
    for i, (idx, valor_banco) in enumerate(opcoes):
        texto_rb = f"Página {idx + 1} do Banco  --->  Valor: R$ {valor_banco}"
        ctk.CTkRadioButton(frame_opcoes, text=texto_rb, variable=var_opcao, value=i, font=("Arial", 15)).pack(pady=15, padx=15, anchor="w")
        
    def confirmar():
        global escolha_global
        escolha_global = var_opcao.get()
        evento_espera.set()
        janela.destroy()
        
    def ignorar():
        global escolha_global
        escolha_global = -1
        evento_espera.set()
        janela.destroy()
        
    janela.protocol("WM_DELETE_WINDOW", ignorar)
    
    frame_btn = ctk.CTkFrame(frame_dir, fg_color="transparent")
    frame_btn.pack(pady=30)
    
    ctk.CTkButton(frame_btn, text="✔️ Salvar Selecionado", command=confirmar, fg_color="green", hover_color="darkgreen", height=50, font=("Arial", 14, "bold")).grid(row=0, column=0, padx=15)
    ctk.CTkButton(frame_btn, text="❌ Pular", command=ignorar, fg_color="red", hover_color="darkred", height=50, font=("Arial", 14, "bold")).grid(row=0, column=1, padx=15)


class AppSeparadorIndependente(ctk.CTk):
    def __init__(self):
        super().__init__()
        
        self.title("Utilitário Independente - Automação de PDFs RH")
        self.geometry("850x650")
        self.state("zoomed")

        titulo = ctk.CTkLabel(self, text="Separador Inteligente de Documentos RH", font=("Arial", 26, "bold"))
        titulo.pack(pady=20)

        frame_botoes = ctk.CTkFrame(self, fg_color="transparent")
        frame_botoes.pack(pady=10)

        botoes_config = [
            ("Separar Contracheques (Via Excel)", lambda: self.acionar_separacao_com_excel("CONTRA CHEQUES"), 0, 0),
            ("Separar Adiantamentos (Via Excel)", lambda: self.acionar_separacao_com_excel("CONTRA ADIANTAMENTOS"), 0, 1),
            ("Separar Pontos (Via Excel)", lambda: self.acionar_separacao_com_excel("PONTOS"), 1, 0),
            ("Comprovantes (Salários)", lambda: self.acionar_comprovantes("COMPROVANTES SALÁRIOS"), 1, 1),
            ("Comprovantes (Adiantamentos)", lambda: self.acionar_comprovantes("COMPROVANTES ADIANTAMENTOS"), 2, 0),
        ]

        for texto, cmd, r, c in botoes_config:
            ctk.CTkButton(frame_botoes, text=texto, command=cmd, width=280, height=45, font=("Arial", 13, "bold")).grid(row=r, column=c, padx=15, pady=12)

        label_console = ctk.CTkLabel(self, text="Console de Processamento:", font=("Arial", 14, "bold"))
        label_console.pack(anchor="w", padx=30, pady=(10, 0))

        self.caixa_texto = ctk.CTkTextbox(
            self, width=800, height=280,
            fg_color="#1E1E1E", text_color="#00FF00", font=("Consolas", 12)
        )
        self.caixa_texto.pack(pady=10, fill="both", expand=True, padx=30)

        self.redirecionador = RedirecionadorConsole(self.caixa_texto)

    def executar_separacao_com_planilha(self, caminho_excel, caminho_pdf, pasta_saida_base, titulo):
        old_stdout = sys.stdout
        sys.stdout = self.redirecionador

        try:
            print(f"\n--- INICIANDO: {titulo} ---")
            
            print("1. A ler nomes da planilha Excel...")
            nomes_brutos = carregar_nomes_de_planilha(caminho_excel)
            nomes_funcionarios = [" ".join(n.strip().upper().split()) for n in nomes_brutos if len(n.strip()) > 2]
            
            if not nomes_funcionarios:
                print("[ERRO] Nenhum nome válido encontrado na planilha Excel.")
                return

            print(f"   -> Foram encontrados {len(nomes_funcionarios)} nomes na planilha.")
            
            nome_pasta_nova = f"Resultado_{titulo.replace(' ', '_')}_{time.strftime('%Y%m%d_%H%M%S')}"
            pasta_saida = os.path.join(pasta_saida_base, nome_pasta_nova)
            os.makedirs(pasta_saida, exist_ok=True)
            print(f"Pasta criada para os ficheiros: {pasta_saida}")

            print("2. A abrir o PDF e a separar (com OCR portátil de emergência se necessário)...")
            try:
                leitor_pdf = PdfReader(caminho_pdf)
            except Exception as e:
                print("Erro ao ler o PDF:", e)
                return

            paginas_salvas = 0
            funcionarios_encontrados = set()

            for numero_pagina, pagina in enumerate(leitor_pdf.pages):
                texto_pagina = extrair_texto_ou_ocr(pagina, caminho_pdf, numero_pagina)
                
                if not texto_pagina:
                    continue

                texto_limpo = " ".join(texto_pagina.upper().split())
                achou_dono = False

                for nome in nomes_funcionarios:
                    if nome[:25].strip() in texto_limpo:
                        escritor_pdf = PdfWriter()
                        escritor_pdf.add_page(pagina)
                        nome_arquivo = "".join(c for c in nome if c.isalnum() or c in " _-")

                        with open(os.path.join(pasta_saida, f"{nome_arquivo}.pdf"), "wb") as arquivo_saida:
                            escritor_pdf.write(arquivo_saida)

                        achou_dono = True
                        paginas_salvas += 1
                        funcionarios_encontrados.add(nome)
                        print(f"[SUCESSO] Encontrado: {nome} (Página {numero_pagina + 1})")
                        break

            print("\n=== PROCESSO CONCLUÍDO ===")
            print(f"Foram gerados {paginas_salvas} ficheiros na pasta: {pasta_saida}")

            nao_encontrados = [n for n in nomes_funcionarios if n not in funcionarios_encontrados]
            if nao_encontrados:
                print(f"\n[ALERTA] {len(nao_encontrados)} colaborador(es) da planilha NÃO foram encontrados no PDF:")
                for n in nao_encontrados:
                    print(f"   - {n}")
            else:
                print("\n[SUCESSO] Todos os colaboradores da planilha foram encontrados no PDF!")
        finally:
            sys.stdout = old_stdout

    def _processar_comprovantes(self, arquivo_gigante, pasta_recibos, pasta_saida_base, titulo, verificar_conflito=False):
        global escolha_global
        old_stdout = sys.stdout
        sys.stdout = self.redirecionador

        try:
            print(f"\n--- INICIANDO: {titulo} ---")
            nome_pasta_nova = f"Resultado_{titulo.replace(' ', '_')}_{time.strftime('%Y%m%d_%H%M%S')}"
            pasta_saida = os.path.join(pasta_saida_base, nome_pasta_nova)
            os.makedirs(pasta_saida, exist_ok=True)
            print(f"Pasta criada para os ficheiros: {pasta_saida}")

            print("1. A ler o PDF gigante do banco e a mapear páginas (com OCR portátil)...")
            try:
                reader_gigante = PdfReader(arquivo_gigante)
            except Exception:
                print("ERRO: Não foi possível abrir o PDF gigante.")
                return

            paginas_comprovantes = []
            for i, pagina in enumerate(reader_gigante.pages):
                texto_pag = normalizar(extrair_texto_ou_ocr(pagina, arquivo_gigante, i))
                paginas_comprovantes.append({
                    "index": i,
                    "texto": texto_pag,
                    "texto_sem_ponto": texto_pag.replace(".", ""),
                    "obj": pagina
                })

            print("2. A ler os recibos individuais e a cruzar dados...\n")
            nao_encontrados = []
            arquivos_gerados = 0

            for arquivo in os.listdir(pasta_recibos):
                if not arquivo.lower().endswith(".pdf"):
                    continue

                caminho_recibo = os.path.join(pasta_recibos, arquivo)
                nome_arquivo_original = normalizar(os.path.splitext(arquivo)[0])
                nome_limpo = re.sub(r"^[\d\s\-]+", "", nome_arquivo_original).strip()
                nome_curto = nome_limpo[:22]

                valores_recibo = extrair_todos_valores(caminho_recibo)
                paginas_matches = [item for item in paginas_comprovantes if nome_curto in item["texto"]]

                paginas_com_valor = []
                for item in paginas_matches:
                    for v in valores_recibo:
                        if v.replace(".", "") in item["texto_sem_ponto"]:
                            paginas_com_valor.append((item, v))
                            break

                if not paginas_com_valor:
                    nao_encontrados.append(nome_arquivo_original)
                    continue

                valor_adiantamento = extrair_valor_linha_adiantamento(caminho_recibo) if verificar_conflito else None
                primeiro_valor_encontrado = paginas_com_valor[0][1]

                if verificar_conflito and valor_adiantamento and primeiro_valor_encontrado == valor_adiantamento:
                    opcoes_para_mostrar = [(p["index"], val) for p, val in paginas_com_valor]
                    escolha_global = -1
                    evento_espera.clear()

                    self.after(0, abrir_janela_escolha, nome_limpo, opcoes_para_mostrar, caminho_recibo, self)
                    print(f"[PAUSA] O sistema esbarrou no adiantamento de: {nome_limpo} (Janela de Escolha)...")
                    evento_espera.wait()

                    if escolha_global != -1:
                        item_escolhido = paginas_com_valor[escolha_global][0]
                        self._salvar_pagina_pdf(item_escolhido["obj"], pasta_saida, nome_arquivo_original)
                        print(f"[SUCESSO] Salvo manualmente: {nome_limpo}")
                        arquivos_gerados += 1
                    else:
                        print("[AVISO] Funcionário ignorado pelo utilizador.")
                        nao_encontrados.append(nome_arquivo_original)
                else:
                    item_escolhido = paginas_com_valor[0][0]
                    self._salvar_pagina_pdf(item_escolhido["obj"], pasta_saida, nome_arquivo_original)
                    print(f"[SUCESSO] Salvo automaticamente: {nome_limpo}")
                    arquivos_gerados += 1

            print("\n=== PROCESSO CONCLUÍDO ===")
            print(f"Foram gerados {arquivos_gerados} ficheiros na pasta: {pasta_saida}")

            if nao_encontrados:
                print(f"\n[ALERTA] {len(nao_encontrados)} recibos sem comprovativos exatos encontrados:")
                for nome in nao_encontrados:
                    print(f"   - {nome}")
            else:
                print("\n[SUCESSO] Todos os recibos tiveram comprovativos encontrados!")
        finally:
            sys.stdout = old_stdout

    @staticmethod
    def _salvar_pagina_pdf(pagina_obj, pasta_saida, nome_arquivo_original):
        writer = PdfWriter()
        writer.add_page(pagina_obj)
        nome_final = f"{nome_arquivo_original} - COMPROVANTE.pdf"
        with open(os.path.join(pasta_saida, nome_final), "wb") as f:
            writer.write(f)

    def acionar_separacao_com_excel(self, titulo):
        caminho_excel = filedialog.askopenfilename(title=f"1. Selecione a Planilha Excel com os Nomes - {titulo}", filetypes=[("Planilhas Excel/CSV", "*.xlsx *.csv")])
        if not caminho_excel:
            return
        pdf_path = filedialog.askopenfilename(title=f"2. Selecione o PDF Geral - {titulo}", filetypes=[("Arquivos PDF", "*.pdf")])
        if not pdf_path:
            return
        saida_path = filedialog.askdirectory(title=f"3. Selecione a Pasta Pai de Resultados - {titulo}")
        if not saida_path:
            return
            
        threading.Thread(target=self.executar_separacao_com_planilha, args=(caminho_excel, pdf_path, saida_path, titulo), daemon=True).start()

    def acionar_comprovantes(self, titulo):
        pdf_gigante = filedialog.askopenfilename(title=f"1. Selecione o PDF Gigante do Banco - {titulo}", filetypes=[("Arquivos PDF", "*.pdf")])
        if not pdf_gigante:
            return
        pasta_recibos = filedialog.askdirectory(title=f"2. Selecione a Pasta com os Recibos Individuais - {titulo}")
        if not pasta_recibos:
            return
        saida_path = filedialog.askdirectory(title=f"3. Selecione a Pasta Pai de Resultados - {titulo}")
        if not saida_path:
            return

        alvo = self._processar_comprovantes
        conflito = True if "SALÁRIOS" in titulo else False
        threading.Thread(target=alvo, args=(pdf_gigante, pasta_recibos, saida_path, titulo, conflito), daemon=True).start()


if __name__ == "__main__":
    ctk.set_appearance_mode("dark")
    ctk.set_default_color_theme("blue")
    
    app = AppSeparadorIndependente()
    app.mainloop()